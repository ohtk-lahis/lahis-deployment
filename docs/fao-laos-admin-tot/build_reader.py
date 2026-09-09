"""Build the LAHIS Admin ToT reader from BOOK.md and chapter Markdown."""
from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
BOOK = ROOT / "BOOK.md"
ACCENT = "155E63"
LIGHT = "E8F2F2"
BORDER = "D9D9D9"
FORBIDDEN_READER_TEXT = (
    "## ลำดับเนื้อหาของเล่ม",
    "## ภาพประกอบที่ใช้สอน",
    "สคริปต์อบรม",
    "สคริปต์เล่า",
    "ต้องเก็บเพิ่ม",
    "ต้องเก็บแบบ crop",
    "รอหน้าจอจริงที่รองรับ",
)


def manifest():
    rows = []
    for line in BOOK.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|") or "---" in line or "Chapter" in line:
            continue
        cells = [x.strip().strip("`") for x in line.strip("|").split("|")]
        if len(cells) >= 4:
            rows.append({"chapter": cells[0], "source": cells[1], "title": cells[2], "audience": cells[3]})
    return rows


def strip_front_matter(text):
    if text.startswith("---\n"):
        _, _, text = text.split("---\n", 2)
    return text.strip() + "\n"


def assemble(rows, output):
    parts = ["# LAHIS Admin Train-the-Trainer Guide\n", "## สารบัญ\n"]
    parts.extend(f"- {r['chapter']} {r['title']}\n" for r in rows)
    for r in rows:
        parts.append("\n\\pagebreak\n\n")
        parts.append(strip_front_matter((ROOT / r["source"]).read_text(encoding="utf-8")))
    output.write_text("".join(parts), encoding="utf-8")


def set_cell_border(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), BORDER)
        borders.append(node)


def apply_font(run, name="Sarabun", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn("w:" + slot), name)
    for attr in list(rfonts.attrib):
        if attr.endswith("Theme"):
            del rfonts.attrib[attr]
    if size is not None:
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        rpr.append(szcs)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def inline(paragraph, text, *, size=None, italic=False, color=None):
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            apply_font(paragraph.add_run(text[pos:match.start()]), size=size, italic=italic, color=color)
        token = match.group(0)
        if token.startswith("`"):
            apply_font(paragraph.add_run(token[1:-1]), "Menlo", size=(size or 13) - 1, color=color)
        else:
            apply_font(paragraph.add_run(token[2:-2]), size=size, bold=True, italic=italic, color=color)
        pos = match.end()
    if pos < len(text):
        apply_font(paragraph.add_run(text[pos:]), size=size, italic=italic, color=color)


def setup_styles(doc):
    specs = {
        "Normal": ("Sarabun", 13, False, "000000", 0, 6),
        "Title": ("Sarabun", 26, True, ACCENT, 0, 12),
        "Subtitle": ("Sarabun", 16, False, "3F5F60", 0, 8),
        "Heading 1": ("Sarabun", 19, True, ACCENT, 10, 8),
        "Heading 2": ("Sarabun", 16, True, ACCENT, 12, 6),
        "Heading 3": ("Sarabun", 14, True, "274748", 10, 4),
        "Caption": ("Sarabun", 10.5, False, "596869", 3, 9),
    }
    for name, (font, size, bold, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for slot in ("ascii", "hAnsi", "cs", "eastAsia"):
            rfonts.set(qn("w:" + slot), font)
        for attr in list(rfonts.attrib):
            if attr.endswith("Theme"):
                del rfonts.attrib[attr]
        style.paragraph_format.line_spacing = 1.18
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_together = True
        style.paragraph_format.keep_with_next = name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3")


def add_numbering(doc, bullet=False):
    root = doc.part.numbering_part.element
    aid = max([int(x.get(qn("w:abstractNumId"))) for x in root.findall(qn("w:abstractNum"))] + [-1]) + 1
    nid = max([int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))] + [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(aid))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("start", "1"), ("numFmt", "bullet" if bullet else "decimal"), ("lvlText", "•" if bullet else "%1."), ("lvlJc", "left")):
        node = OxmlElement("w:" + tag)
        node.set(qn("w:val"), val)
        lvl.append(node)
    abstract.append(lvl)
    first_num = root.find(qn("w:num"))
    root.insert(list(root).index(first_num) if first_num is not None else len(root), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(nid))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(aid))
    num.append(ref)
    root.append(num)
    return nid


def set_list(paragraph, num_id):
    pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    pr.get_or_add_ilvl().val = 0
    pr.get_or_add_numId().val = num_id
    paragraph.paragraph_format.left_indent = Inches(0.32)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.paragraph_format.space_after = Pt(3)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    apply_font(run, size=9, color="6A7373")


def convert_svg(path, cache):
    target = cache / (path.name + ".png")
    if not target.exists():
        # Quick Look renders SVG thumbnails into a square and scales from the
        # shorter side. A wide diagram is therefore clipped on the right. Give
        # it a temporary square viewBox, then crop the resulting PNG back to
        # the SVG's declared aspect ratio.
        root = ET.parse(path).getroot()
        viewbox = root.get("viewBox", "").split()
        if len(viewbox) != 4:
            raise ValueError(f"SVG has no usable viewBox: {path}")
        x, y, width, height = map(float, viewbox)
        side = max(width, height)
        padded = cache / (path.stem + "-padded.svg")
        root.set("width", str(side))
        root.set("height", str(side))
        root.set("viewBox", f"{x:g} {y:g} {side:g} {side:g}")
        ET.ElementTree(root).write(padded, encoding="utf-8", xml_declaration=True)
        subprocess.run(
            ["qlmanage", "-t", "-s", "1800", "-o", str(cache), str(padded)],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        square = cache / (padded.name + ".png")
        with Image.open(square) as im:
            crop_w = round(im.width * width / side)
            crop_h = round(im.height * height / side)
            im.crop((0, 0, crop_w, crop_h)).save(target)
    return target


def add_image(doc, path, alt, cache):
    if path.suffix.lower() == ".svg":
        path = convert_svg(path, cache)
    with Image.open(path) as im:
        w, h = im.size
    max_w, max_h = 6.35, 6.15
    ratio = min(max_w / w, max_h / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(w * ratio), height=Inches(h * ratio))
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        drawing[0].set("descr", alt)


def add_table(doc, rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.autofit = True
    table.style = "Table Grid"
    for ri, data in enumerate(rows):
        cells = table.add_row().cells
        trpr = table.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for ci, (cell, value) in enumerate(zip(cells, data)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            tcpr = cell._tc.get_or_add_tcPr()
            margin = OxmlElement("w:tcMar")
            for side, val in (("top", 70), ("bottom", 70), ("left", 90), ("right", 90)):
                node = OxmlElement("w:" + side)
                node.set(qn("w:w"), str(val))
                node.set(qn("w:type"), "dxa")
                margin.append(node)
            tcpr.append(margin)
            if ri == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), ACCENT)
                tcpr.append(shd)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            inline(p, value, size=10.5, color="FFFFFF" if ri == 0 else "000000")
            for run in p.runs:
                run.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_blocks(text):
    lines = strip_front_matter(text).splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line == "\\pagebreak":
            blocks.append(("pagebreak", None)); i += 1; continue
        if line.startswith("```"):
            lang = line[3:].strip(); content = []; i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                content.append(lines[i]); i += 1
            i += 1
            blocks.append(("code", (lang, "\n".join(content))))
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                row = [x.strip() for x in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", x) for x in row):
                    rows.append(row)
                i += 1
            blocks.append(("table", rows)); continue
        image = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if image:
            blocks.append(("image", image.groups())); i += 1; continue
        head = re.match(r"^(#{1,3})\s+(.+)$", line)
        if head:
            blocks.append(("heading", (len(head.group(1)), head.group(2)))); i += 1; continue
        ordered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if ordered:
            items = []
            while i < len(lines) and (m := re.match(r"^(\d+)\.\s+(.+)$", lines[i])):
                items.append(m.group(2)); i += 1
            blocks.append(("ordered", items)); continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            items = []
            while i < len(lines) and (m := re.match(r"^-\s+(.+)$", lines[i])):
                items.append(m.group(1)); i += 1
            blocks.append(("bullet", items)); continue
        if line.startswith("*ภาพ") and line.endswith("*"):
            blocks.append(("caption", line[1:-1])); i += 1; continue
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3})\s|^```|^\||^!\[|^\d+\.\s|^-\s|^\*ภาพ|^\\pagebreak$", lines[i]):
            para.append(lines[i]); i += 1
        blocks.append(("paragraph", " ".join(x.strip() for x in para)))
    return blocks


def build(rows, output):
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.78)
    page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline(p, "LAHIS Admin Train-the-Trainer Guide", size=26, color=ACCENT)
    p.paragraph_format.space_before = Inches(2.0)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline(p, "คู่มือผู้ดูแลระบบและผู้ฝึกสอน สำหรับการอบรมใน สปป.ลาว", size=16, color="3F5F60")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    inline(p, "FAO Laos · LAHIS", size=12, color="6A7373")

    doc.add_page_break()
    doc.add_heading("สารบัญ", level=1)
    toc = doc.add_table(rows=0, cols=2)
    toc.autofit = False
    for r in rows:
        cells = toc.add_row().cells
        cells[0].width, cells[1].width = Inches(0.65), Inches(5.8)
        for cell in cells:
            tcpr = cell._tc.get_or_add_tcPr()
            margin = OxmlElement("w:tcMar")
            for side, val in (("top", 45), ("bottom", 45), ("left", 20), ("right", 20)):
                node = OxmlElement("w:" + side); node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa"); margin.append(node)
            tcpr.append(margin)
        inline(cells[0].paragraphs[0], r["chapter"], size=12, color=ACCENT)
        inline(cells[1].paragraphs[0], r["title"], size=12)

    with tempfile.TemporaryDirectory(prefix="lahis-svg-") as temp:
        cache = Path(temp)
        for r in rows:
            doc.add_page_break()
            blocks = parse_blocks((ROOT / r["source"]).read_text(encoding="utf-8"))
            for bi, (kind, value) in enumerate(blocks):
                if kind == "heading":
                    level, text = value
                    p = doc.add_paragraph(style=f"Heading {level}")
                    inline(p, text)
                elif kind == "paragraph":
                    p = doc.add_paragraph()
                    inline(p, value)
                    if bi + 1 < len(blocks) and blocks[bi + 1][0] in ("image", "table"):
                        p.paragraph_format.keep_with_next = True
                elif kind in ("ordered", "bullet"):
                    num = add_numbering(doc, kind == "bullet")
                    for item in value:
                        p = doc.add_paragraph()
                        inline(p, item)
                        set_list(p, num)
                elif kind == "image":
                    alt, rel = value
                    add_image(doc, (ROOT / "chapters" / rel).resolve(), alt, cache)
                elif kind == "caption":
                    p = doc.add_paragraph(style="Caption")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    inline(p, value, size=10.5, italic=True, color="596869")
                elif kind == "table":
                    add_table(doc, value)
                elif kind == "code":
                    lang, code = value
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.25)
                    p.paragraph_format.right_indent = Inches(0.25)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(8)
                    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F2F4F4"); p._p.get_or_add_pPr().append(shd)
                    apply_font(p.add_run(code), "Menlo", 9.5, color="263333")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", type=Path, default=ROOT / "FAO_LAOS_ADMIN_TOT.md")
    parser.add_argument("--docx", type=Path, default=ROOT / "FAO_LAOS_ADMIN_TOT.docx")
    args = parser.parse_args()
    rows = manifest()
    if len(rows) != 16:
        raise SystemExit(f"Expected 16 manifest rows, found {len(rows)}")
    for row in rows:
        path = ROOT / row["source"]
        text = path.read_text(encoding="utf-8")
        for forbidden in FORBIDDEN_READER_TEXT:
            if forbidden in strip_front_matter(text):
                raise SystemExit(f"Editorial text leaked into reader content: {path}: {forbidden!r}")
        h1 = re.search(r"^# (.+)$", text, re.M)
        expected = f"{row['chapter']} {row['title']}"
        if not h1 or h1.group(1) != expected:
            raise SystemExit(f"Manifest mismatch: {path}: expected {expected!r}")
        for _, rel in re.findall(r"!\[(.*?)\]\((.*?)\)", text):
            if not (path.parent / rel).resolve().exists():
                raise SystemExit(f"Missing image: {path.parent / rel}")
    assemble(rows, args.markdown)
    build(rows, args.docx)
    print(args.markdown)
    print(args.docx)


if __name__ == "__main__":
    main()
