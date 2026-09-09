"""Small reproducible probes. No publication or production-document mutation."""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

parser = argparse.ArgumentParser()
parser.add_argument('--source', type=Path, default=Path(__file__).with_name('01-thai.md'))
parser.add_argument('--output', type=Path, required=True)
args = parser.parse_args()
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.75)
for name, size in [('Normal', 12), ('Title', 20), ('Heading 1', 16)]:
    style = doc.styles[name]
    style.font.name = 'Tahoma'
    style.font.size = Pt(size)
    for slot in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        style.element.get_or_add_rPr().rFonts.set(qn('w:' + slot), 'Tahoma')
    style.paragraph_format.line_spacing = 1.2
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.space_before = Pt(12 if name == 'Heading 1' else 0)
    style.paragraph_format.keep_with_next = name != 'Normal'
for line in args.source.read_text().splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], 'Heading 1')
    else:
        doc.add_paragraph(line)
args.output.parent.mkdir(parents=True, exist_ok=True)
doc.save(args.output)
print(args.output)
