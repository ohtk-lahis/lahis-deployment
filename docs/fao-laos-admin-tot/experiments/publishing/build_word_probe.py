"""Build the small Markdown fixture as a Word document; no cloud mutations."""
import argparse
import json
import subprocess
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

p=argparse.ArgumentParser()
p.add_argument('--source',type=Path,default=Path(__file__).with_name('03-layout.md'))
p.add_argument('--output',type=Path,required=True)
a=p.parse_args()
model=json.loads(subprocess.check_output([sys.executable,str(Path(__file__).with_name('compile_native_probe.py')),str(a.source)]))
root=Path(__file__).resolve().parents[2]
# Two existing lesson crops; aliases keep the text fixture unchanged.
images={'i.0':root/'screenshots/c06a-02-definition-form-builder.jpg',
        'i.11':root/'screenshots/c09-03b-pending-offline.png'}
doc=Document()
s=doc.sections[0]
s.page_width=Pt(595.28);s.page_height=Pt(841.89)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Pt(54)
for name,size in [('Normal',12),('Title',20),('Heading 1',15),('Caption',10)]:
    st=doc.styles[name];st.font.name='Sarabun';st.font.size=Pt(size)
    for slot in ('ascii','hAnsi','cs','eastAsia'):
        st.element.get_or_add_rPr().rFonts.set(qn('w:'+slot),'Sarabun')
    for attr in list(st.element.rPr.rFonts.attrib):
        if attr.endswith('Theme'):del st.element.rPr.rFonts.attrib[attr]
    cs=OxmlElement('w:szCs');cs.set(qn('w:val'),str(size*2));st.element.rPr.append(cs)
    from docx.shared import RGBColor
    st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.line_spacing=1.2
    st.paragraph_format.space_before=Pt(12 if name=='Heading 1' else 0)
    st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.keep_together=True
    st.paragraph_format.keep_with_next=name in ('Title','Heading 1')
    st.font.bold=name in ('Title','Heading 1')

def numbering(bullet=False):
    numbering=doc.part.numbering_part.element
    aid=max([int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]+[-1])+1
    nid=max([int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]+[0])+1
    abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),str(aid))
    lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
    for tag,val in [('start','1'),('numFmt','bullet' if bullet else 'decimal'),('lvlText','•' if bullet else '%1.'),('lvlJc','left')]:
        n=OxmlElement('w:'+tag);n.set(qn('w:val'),val);lvl.append(n)
    abstract.append(lvl)
    first_num=numbering.find(qn('w:num'))
    numbering.insert(list(numbering).index(first_num) if first_num is not None else len(numbering),abstract)
    num=OxmlElement('w:num');num.set(qn('w:numId'),str(nid))
    n=OxmlElement('w:abstractNumId');n.set(qn('w:val'),str(aid));num.append(n)
    override=OxmlElement('w:lvlOverride');override.set(qn('w:ilvl'),'0')
    start=OxmlElement('w:startOverride');start.set(qn('w:val'),'1');override.append(start);num.append(override)
    numbering.append(num);return nid

groups={}
for i,b in enumerate(model['blocks']):
    kind=b['kind']
    if kind=='table':
        t=doc.add_table(rows=0,cols=len(b['rows'][0]));t.style='Table Grid';t.autofit=False
        widths=[56,122,309.28]
        for c,w in zip(t.columns,widths):c.width=Pt(w)
        for ri,row in enumerate(b['rows']):
            cells=t.add_row().cells
            trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if ri==0:trpr.append(OxmlElement('w:tblHeader'))
            for ci,(cell,text) in enumerate(zip(cells,row)):
                cell.width=Pt(widths[ci]);cell.text=text
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(0)
                    para.paragraph_format.line_spacing=1.1
                    para.paragraph_format.keep_with_next=ri==0
                    for run in para.runs:run.bold=ri==0
                tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
                for side,v in [('top',80),('bottom',80),('left',120),('right',120)]:
                    n=OxmlElement('w:'+side);n.set(qn('w:w'),str(v));n.set(qn('w:type'),'dxa');m.append(n)
                tcpr.append(m)
                if ri==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EFF5');tcpr.append(fill)
    elif kind=='image':
        path=images[b['imageId']]
        with Image.open(path) as im:w,h=im.size
        scale=min(432/w,288/h)
        para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next=True
        para.add_run().add_picture(str(path),width=Pt(w*scale),height=Pt(h*scale))
    else:
        para=doc.add_paragraph(b['text'],{'title':'Title','heading':'Heading 1','caption':'Caption'}.get(kind,'Normal'))
        if kind=='caption':para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if i+1<len(model['blocks']) and model['blocks'][i+1]['kind'] in ('image','table'):para.paragraph_format.keep_with_next=True
        if kind in ('number','bullet'):
            group=('number',b['group']) if kind=='number' else ('bullet',0)
            if group not in groups:groups[group]=numbering(kind=='bullet')
            pr=para._p.get_or_add_pPr().get_or_add_numPr()
            pr.get_or_add_ilvl().val=0;pr.get_or_add_numId().val=groups[group]
            para.paragraph_format.left_indent=Pt(24);para.paragraph_format.first_line_indent=Pt(-16)
            para.paragraph_format.space_after=Pt(3)
a.output.parent.mkdir(parents=True,exist_ok=True)
doc.save(a.output)
print(a.output)
